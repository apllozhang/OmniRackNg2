from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.section import WD_ORIENT
from PIL import Image

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, Table, TableStyle, PageBreak, KeepTogether
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT.parent.parent / 'outputs'
OUT.mkdir(parents=True, exist_ok=True)
ASSET = ROOT / 'assets' / 'product-doc'
TODAY = '2026年9月4日'
PRODUCT_VERSION = '0.1.0'
DOC_VERSION = '1.0.0'
DOCX_PATH = OUT / f'机柜管理工具-产品宣传文档-{PRODUCT_VERSION}.docx'
PDF_PATH = OUT / f'机柜管理工具-产品宣传文档-{PRODUCT_VERSION}.pdf'
MD_PATH = OUT / f'机柜管理工具-产品宣传文档-{PRODUCT_VERSION}.md'
PROJECT_MD_PATH = ROOT / 'docs' / f'产品宣传文档-{PRODUCT_VERSION}.md'

SCREENSHOTS = {
    'resource': ASSET / '01-resource-hierarchy.png',
    'screen_menu': ASSET / '02-big-screen-menu.png',
    'rack_detail': ASSET / '03-rack-detail-capacity.png',
    'search': ASSET / '04-search-location.png',
    'device_editor': ASSET / '05-device-editor.png',
    'room_status': ASSET / '06-room-screen-status.png',
    'batch_editor': ASSET / '07-batch-editor.png',
}

# ---------- shared content ----------

def md_content():
    return f'''# 机柜管理工具
## 数据中心基础资源可视化管理平台

> **让每一台设备、每一个 U 位、每一次变更都清晰可见。**  
> 产品版本：**{PRODUCT_VERSION}** ｜ 文档版本：**{DOC_VERSION}** ｜ 更新日期：**{TODAY}**

![机房大屏总览](../work/project-current/assets/product-doc/02-big-screen-menu.png)

## 1. 产品定位

机柜管理工具面向数据中心、机房、机柜和设备的日常管理，将资源台账、U 位布局、设备信息、容量分析和运维操作统一在一个平台中。系统以机房大屏为核心入口，以结构化数据为基础，帮助运维团队从“靠记忆、靠表格、靠现场查找”转向“可视、可查、可操作、可追溯”。

平台适合新建机房初始化、存量资源盘点、设备批量上架、设备迁移、机柜扩容、运维交接和多机房统一管理等场景。

## 2. 核心价值

- **资源全景可视化**：从数据中心到 U 位建立统一资源视图，机房状态一眼可见。
- **设备定位精确到 U 位**：通过设备名称搜索，快速定位设备所在机柜和 U 位，减少现场查找时间。
- **变更操作形成闭环**：设备上架、编辑、迁移、下架和归档均由业务流程管理，位置变化可追溯。
- **批量维护更高效**：提供机柜、设备和机柜图的 Excel 模板、导入、导出及回导能力。
- **容量分析更可信**：对 U 位、功耗、承重和 PDU 供电分别计算；条件不足时明确给出无法计算原因。
- **跨资源域调整更简单**：支持数据中心、机房和机柜复制，以及跨机房、跨数据中心移动资源。
- **权限管理更稳妥**：支持本地用户、角色权限、管理员审批和 LDAP 对接，满足团队化运维要求。

## 3. 产品能力总览

| 能力域 | 主要能力 |
|---|---|
| 资源管理 | 数据中心、机房、机柜的增删改查、状态维护、复制与层级浏览 |
| 机房大屏 | 多机柜网格展示、固定格位、拖动交换位置、动态扩展和换行 |
| U 位管理 | 空 U 位上架、有设备 U 位查看/编辑/迁移/下架，统一左键操作 |
| 设备台账 | 设备信息维护、自动编码、统一编辑、搜索定位、位置履历 |
| 模板与批量 | 机柜模板、机柜/设备批量导入、模板下载、设备信息导出后回导更新 |
| 机柜图 | 机柜图 Excel 导出与回导，支持导入验证和冲突检查 |
| 容量与供电 | U 位、功耗、承重、PDU 插座与供电连接关系分析 |
| 权限治理 | 用户管理、角色权限、管理员审批、LDAP 用户和组映射、操作审计 |

## 4. 重点功能展示

### 4.1 统一资源层级管理

系统按照“数据中心 → 机房 → 机柜 → U 位 → 设备”组织资源。管理人员可以在同一页面浏览层级关系、查看资源数量和状态，并直接新增、编辑或进入更多操作。

![统一资源层级管理](../work/project-current/assets/product-doc/01-resource-hierarchy.png)

**使用价值**：避免数据中心、机房和机柜信息分散在不同表格中，建立统一资源主数据。

### 4.2 机房大屏与机柜可视化

机房大屏以固定格位展示机柜，机柜不会相互重叠；支持通过拖动交换位置，并按照机房配置自动扩展和换行。管理人员可在大屏上直接新增、编辑、删除机柜和修改机柜状态。

![机房大屏总览](../work/project-current/assets/product-doc/02-big-screen-menu.png)

**使用价值**：将“机柜在哪里、空了多少、用了多少、状态如何”从表格信息变成直观空间视图。

### 4.3 U 位操作与设备生命周期

平台统一使用左键点击 U 位：空 U 位可直接上架设备，有设备的 U 位可查看详情、编辑设备信息、调整/迁移位置或执行下架。设备位置由后端业务规则校验，避免重叠和重复占用。

![设备 U 位操作菜单](../work/project-current/assets/product-doc/02-big-screen-menu.png)

### 4.4 设备搜索与快速定位

输入设备名称并点击查询，即可在大屏中快速定位目标设备所在机柜及 U 位，适合设备巡检、故障处理和变更确认。

![设备搜索定位](../work/project-current/assets/product-doc/04-search-location.png)

### 4.5 详情与容量分析

机柜详情左侧展示机柜图，右侧集中展示机柜参数、U 位利用率、功耗利用率、承重利用率和 PDU 端口利用率。设备或机柜未补充必要数据时，系统不将缺失值默认为 0，而是输出无法计算的具体原因。

![机柜详情与容量分析](../work/project-current/assets/product-doc/03-rack-detail-capacity.png)

### 4.6 统一设备编辑

设备台账和机房大屏共用同一套设备编辑表单，字段以后台数据模型为准，确保新增、编辑、详情、批量导入和导出回导之间的数据一致。

支持设备名称、设备编号、类型、厂商、型号、序列号、管理 IP、业务 IP、功耗、U 高度、负责人、用途和备注等信息；设备编号支持自动生成，也支持按管理规范手工维护。

![统一设备编辑表单](../work/project-current/assets/product-doc/05-device-editor.png)

### 4.7 模板、批量导入与设备信息回导

平台提供清晰的 Excel 模板，区分必填字段和可选字段，并提供示例、字段说明和校验提示。普通机柜和设备批量导入不要求填写编码，编码为空时由系统自动生成。

设备信息可先导出，再补充管理 IP、业务 IP、SN、功耗等字段后重新导入更新。机柜图也支持导出和回导，导入阶段先执行格式、设备识别、位置占用和重复数据验证，确认后再提交变更，降低数据错乱风险。

![批量信息维护](../work/project-current/assets/product-doc/07-batch-editor.png)

### 4.8 复制与跨机房/跨数据中心移动

支持复制数据中心、机房和机柜，用于快速构造相似资源环境；支持跨机房、跨数据中心移动机房或机柜，并在移动前校验目标层级、名称/编码冲突、U 位占用和设备位置关系。

### 4.9 PDU 与供电数据

平台支持 PDU、插座和设备供电连接关系管理，插座类型可区分国标、欧标、10A 和 16A。PDU 信息属于可选配置，不会阻塞基础资源和设备管理；补充完整数据后，系统可进一步输出供电使用情况和利用率。

## 5. 典型工作流程

1. 创建数据中心，编码可自动生成。
2. 创建机房并设置机房基本信息和每排最多展示的机柜数量。
3. 通过机柜模板创建机柜，按实际情况调整 U 位数、宽度、深度、高度等参数。
4. 新增设备或通过 Excel 批量导入设备，编码可自动生成。
5. 在大屏空 U 位上架设备，或通过迁移流程调整设备位置。
6. 通过设备名称查询，快速定位设备所在机柜和 U 位。
7. 查看机柜详情和利用率；补充功耗、承重或 PDU 数据后获得更完整分析。
8. 进行设备编辑、迁移、下架、复制、跨机房移动并保留操作记录。

## 6. 典型应用场景

### 新建数据中心初始化
统一创建数据中心、机房和机柜，使用模板和批量导入快速完成基础数据初始化。

### 存量机房资源盘点
将现有设备信息整理后导入，使用大屏查看机柜占用情况，通过搜索快速核对设备位置。

### 设备批量变更
导出设备信息，集中补充 IP、SN、功耗和责任人等字段，复核后批量回导更新。

### 机房扩容与资源规划
设置每排机柜数量，拖动调整大屏布局；通过 U 位、功耗、承重和 PDU 利用率辅助判断扩容空间。

### 设备迁移与运维交接
在大屏中直接完成设备编辑、迁移和下架，结合位置履历、审计和审批记录，降低交接遗漏。

### 多机房统一管理
按数据中心和机房分层查看资源，复制相似环境并在资源域之间移动机柜和机房。

## 7. 产品特点与管理保障

- **编码灵活**：数据中心、机房、机柜、设备和模板支持自动生成编码，也支持手工编码。
- **布局可靠**：机柜使用固定格位，拖动只能交换或调整到有效位置，避免视觉和数据层面的重叠。
- **字段一致**：设备台账、大屏编辑、详情、导入和导出以同一后台字段模型为准。
- **全系统统一视图**：不区分设备正面和背面，减少维护歧义。
- **容量分析不虚报**：缺少功耗、承重或 PDU 数据时展示“无法计算”及原因。
- **资源可持续扩展**：不以固定业务数量上限限制数据中心、机房和机柜的持续建设。
- **安全可控**：管理员权限、审批策略、LDAP 组映射和审计能力共同保障关键变更。
- **导入可复核**：批量导入和机柜图回导先验证、后确认、再提交，支持识别新增、更新和潜在冲突。

## 8. 部署形态与技术基础

平台采用前后端分离架构，可通过 Docker Compose 快速部署；数据持久化使用 PostgreSQL。前端采用 Vue 3、TypeScript 和 Element Plus，后端采用 Go、Gin 和 GORM。LDAP 作为可选企业目录服务对接能力，可按组织现有身份体系接入。

## 9. 对外产品信息

| 项目 | 内容 |
|---|---|
| 产品名称 | 机柜管理工具 |
| 产品定位 | 数据中心基础资源可视化管理平台 |
| 正式产品版本 | {PRODUCT_VERSION} |
| 适用对象 | 数据中心运维团队、机房管理员、基础设施管理人员、审计与管理人员 |
| 部署方式 | Docker Compose / 私有化部署 |
| 数据存储 | PostgreSQL |
| 认证方式 | 本地账号、可选 LDAP 对接 |

## 10. 结语

从资源台账到机房大屏，从设备上架到跨机房调整，机柜管理工具让数据中心资源管理更加清晰、准确、高效。平台将每一台设备、每一个 U 位和每一次变更连接起来，为数据中心的日常运维、容量规划和持续扩展提供统一基础。

> **机柜管理工具**  
> 让基础设施数据真正成为可视、可用、可管理的运营资产。

---

> **截图说明**：本文使用项目演示环境截图，截图中的名称、编号和时间为演示数据，仅用于功能说明，不代表任何客户生产环境。
'''

# ---------- docx helpers ----------
def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell(cell, text, bold=False, color=(31,78,121), size=9):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text)); r.bold = bold; r.font.size = Pt(size)
    r.font.name = 'Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    r.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers): set_cell(t.rows[0].cells[i], h, True, (255,255,255), 9); shade_cell(t.rows[0].cells[i], '0F5B78')
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            set_cell(cells[i], val, False, (35,45,55), 8)
            if ridx%2: shade_cell(cells[i], 'EDF6F8')
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def add_bullet(doc, text, level=0):
    p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); p.paragraph_format.space_after=Pt(2); p.add_run(text)

def add_number(doc, text):
    p=doc.add_paragraph(style='List Number'); p.paragraph_format.space_after=Pt(2); p.add_run(text)

def add_caption(doc, text):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(10)
    r=p.add_run(text); r.italic=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor(105,105,105)

def add_image_docx(doc, path, caption, width=6.7):
    if not path.exists(): return
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)

def docx_setup(doc):
    for sec in doc.sections:
        sec.top_margin=Inches(.62); sec.bottom_margin=Inches(.62); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
    styles=doc.styles
    normal=styles['Normal']; normal.font.name='Microsoft YaHei'; normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); normal.font.size=Pt(10); normal.paragraph_format.space_after=Pt(6)
    for name,size,color in [('Title',30,'0F5B78'),('Heading 1',18,'0F5B78'),('Heading 2',13,'127C8C'),('Heading 3',11,'2E7D32')]:
        st=styles[name]; st.font.name='Microsoft YaHei'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color); st.font.bold=True
    for sec in doc.sections:
        footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=footer.add_run(f'机柜管理工具 · 产品宣传文档 · V{PRODUCT_VERSION} · {TODAY}'); r.font.size=Pt(8); r.font.color.rgb=RGBColor(120,120,120)

def build_docx():
    d=Document(); docx_setup(d)
    d.core_properties.title='机柜管理工具产品宣传文档'
    d.core_properties.subject='数据中心基础资源可视化管理平台'
    d.core_properties.author='Cabinet Management Tool'
    # cover
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(70)
    r=p.add_run('机柜管理工具'); r.font.size=Pt(32); r.font.bold=True; r.font.color.rgb=RGBColor(15,91,120)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('数据中心基础资源可视化管理平台'); r.font.size=Pt(18); r.font.color.rgb=RGBColor(46,125,50)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(14); r=p.add_run('让数据中心资源管理更清晰、更准确、更高效'); r.font.size=Pt(13); r.font.color.rgb=RGBColor(90,90,90)
    d.add_paragraph('')
    add_image_docx(d, SCREENSHOTS['screen_menu'], '图 1 机房大屏示例（演示数据）', width=6.8)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(25)
    r=p.add_run(f'正式产品版本：V{PRODUCT_VERSION}\n文档版本：V{DOC_VERSION}　|　更新日期：{TODAY}'); r.bold=True; r.font.size=Pt(10); r.font.color.rgb=RGBColor(90,90,90)
    d.add_page_break()

    d.add_heading('目录', level=1)
    for x in ['1. 产品定位','2. 核心价值','3. 产品能力总览','4. 重点功能展示','5. 典型工作流程','6. 典型应用场景','7. 产品特点与管理保障','8. 部署形态与技术基础','9. 对外产品信息','10. 结语']:
        d.add_paragraph(x, style='List Number')
    d.add_page_break()

    d.add_heading('1. 产品定位', 1)
    d.add_paragraph('机柜管理工具面向数据中心、机房、机柜和设备的日常管理，将资源台账、U 位布局、设备信息、容量分析和运维操作统一在一个平台中。系统以机房大屏为核心入口，以结构化数据为基础，帮助运维团队从“靠记忆、靠表格、靠现场查找”转向“可视、可查、可操作、可追溯”。')
    d.add_paragraph('平台适合新建机房初始化、存量资源盘点、设备批量上架、设备迁移、机柜扩容、运维交接和多机房统一管理等场景。')
    d.add_heading('2. 核心价值', 1)
    for s in ['资源全景可视化：从数据中心到 U 位建立统一资源视图，机房状态一眼可见。','设备定位精确到 U 位：通过设备名称搜索，快速定位设备所在机柜和 U 位，减少现场查找时间。','变更操作形成闭环：设备上架、编辑、迁移、下架和归档均由业务流程管理，位置变化可追溯。','批量维护更高效：提供机柜、设备和机柜图的 Excel 模板、导入、导出及回导能力。','容量分析更可信：对 U 位、功耗、承重和 PDU 供电分别计算；条件不足时明确给出无法计算原因。','跨资源域调整更简单：支持数据中心、机房和机柜复制，以及跨机房、跨数据中心移动资源。','权限管理更稳妥：支持本地用户、角色权限、管理员审批和 LDAP 对接，满足团队化运维要求。']: add_bullet(d,s)
    d.add_heading('3. 产品能力总览', 1)
    table(d,['能力域','主要能力'],[['资源管理','数据中心、机房、机柜的增删改查、状态维护、复制与层级浏览'],['机房大屏','多机柜网格展示、固定格位、拖动交换位置、动态扩展和换行'],['U 位管理','空 U 位上架、有设备 U 位查看/编辑/迁移/下架，统一左键操作'],['设备台账','设备信息维护、自动编码、统一编辑、搜索定位、位置履历'],['模板与批量','机柜模板、机柜/设备批量导入、模板下载、设备信息导出后回导更新'],['机柜图','机柜图 Excel 导出与回导，支持导入验证和冲突检查'],['容量与供电','U 位、功耗、承重、PDU 插座与供电连接关系分析'],['权限治理','用户管理、角色权限、管理员审批、LDAP 用户和组映射、操作审计']],[1.15,5.4])
    d.add_page_break()

    d.add_heading('4. 重点功能展示', 1)
    d.add_heading('4.1 统一资源层级管理', 2)
    d.add_paragraph('系统按照“数据中心 → 机房 → 机柜 → U 位 → 设备”组织资源。管理人员可以在同一页面浏览层级关系、查看资源数量和状态，并直接新增、编辑或进入更多操作。')
    add_image_docx(d, SCREENSHOTS['resource'], '图 2 统一资源层级管理：以层级关系承载资源主数据', 6.8)
    d.add_paragraph('使用价值：避免数据中心、机房和机柜信息分散在不同表格中，建立统一资源主数据。')
    d.add_heading('4.2 机房大屏与机柜可视化', 2)
    d.add_paragraph('机房大屏以固定格位展示机柜，机柜不会相互重叠；支持通过拖动交换位置，并按照机房配置自动扩展和换行。管理人员可在大屏上直接新增、编辑、删除机柜和修改机柜状态。')
    add_image_docx(d, SCREENSHOTS['room_status'], '图 3 机房大屏：多机柜状态、U 位占用和固定格位展示', 6.8)
    d.add_heading('4.3 U 位操作与设备生命周期', 2)
    d.add_paragraph('平台统一使用左键点击 U 位：空 U 位可直接上架设备，有设备的 U 位可查看详情、编辑设备信息、调整/迁移位置或执行下架。设备位置由后端业务规则校验，避免重叠和重复占用。')
    add_image_docx(d, SCREENSHOTS['screen_menu'], '图 4 U 位操作菜单：详情、调整/迁移、下架等操作集中呈现', 6.8)
    d.add_page_break()

    d.add_heading('4.4 设备搜索与快速定位', 2)
    d.add_paragraph('输入设备名称并点击查询，即可在大屏中快速定位目标设备所在机柜及 U 位，适合设备巡检、故障处理和变更确认。')
    add_image_docx(d, SCREENSHOTS['search'], '图 5 设备搜索：从名称快速定位机柜和 U 位', 6.8)
    d.add_heading('4.5 详情与容量分析', 2)
    d.add_paragraph('机柜详情左侧展示机柜图，右侧集中展示机柜参数、U 位利用率、功耗利用率、承重利用率和 PDU 端口利用率。设备或机柜未补充必要数据时，系统不将缺失值默认为 0，而是输出无法计算的具体原因。')
    add_image_docx(d, SCREENSHOTS['rack_detail'], '图 6 机柜详情与容量分析：参数、U 位、功耗和 PDU 信息一体化展示', 6.8)
    d.add_heading('4.6 统一设备编辑', 2)
    d.add_paragraph('设备台账和机房大屏共用同一套设备编辑表单，字段以后台数据模型为准，确保新增、编辑、详情、批量导入和导出回导之间的数据一致。')
    d.add_paragraph('支持设备名称、设备编号、类型、厂商、型号、序列号、管理 IP、业务 IP、功耗、U 高度、负责人、用途和备注等信息；设备编号支持自动生成，也支持按管理规范手工维护。')
    add_image_docx(d, SCREENSHOTS['device_editor'], '图 7 统一设备编辑表单：台账与大屏使用同一字段模型', 6.2)
    d.add_page_break()

    d.add_heading('4.7 模板、批量导入与设备信息回导', 2)
    d.add_paragraph('平台提供清晰的 Excel 模板，区分必填字段和可选字段，并提供示例、字段说明和校验提示。普通机柜和设备批量导入不要求填写编码，编码为空时由系统自动生成。')
    d.add_paragraph('设备信息可先导出，再补充管理 IP、业务 IP、SN、功耗等字段后重新导入更新。机柜图也支持导出和回导，导入阶段先执行格式、设备识别、位置占用和重复数据验证，确认后再提交变更，降低数据错乱风险。')
    add_image_docx(d, SCREENSHOTS['batch_editor'], '图 8 批量维护：通过模板和字段校验降低批量操作风险', 6.8)
    d.add_heading('4.8 复制与跨机房/跨数据中心移动', 2)
    d.add_paragraph('支持复制数据中心、机房和机柜，用于快速构造相似资源环境；支持跨机房、跨数据中心移动机房或机柜，并在移动前校验目标层级、名称/编码冲突、U 位占用和设备位置关系。')
    d.add_heading('4.9 PDU 与供电数据', 2)
    d.add_paragraph('平台支持 PDU、插座和设备供电连接关系管理，插座类型可区分国标、欧标、10A 和 16A。PDU 信息属于可选配置，不会阻塞基础资源和设备管理；补充完整数据后，系统可进一步输出供电使用情况和利用率。')

    d.add_heading('5. 典型工作流程', 1)
    for s in ['创建数据中心，编码可自动生成。','创建机房并设置机房基本信息和每排最多展示的机柜数量。','通过机柜模板创建机柜，按实际情况调整 U 位数、宽度、深度、高度等参数。','新增设备或通过 Excel 批量导入设备，编码可自动生成。','在大屏空 U 位上架设备，或通过迁移流程调整设备位置。','通过设备名称查询，快速定位设备所在机柜和 U 位。','查看机柜详情和利用率；补充功耗、承重或 PDU 数据后获得更完整分析。','进行设备编辑、迁移、下架、复制、跨机房移动并保留操作记录。']: add_number(d,s)
    d.add_heading('6. 典型应用场景', 1)
    table(d,['场景','平台价值'],[['新建数据中心初始化','统一创建数据中心、机房和机柜，使用模板和批量导入快速完成基础数据初始化。'],['存量机房资源盘点','导入现有设备信息，用大屏查看机柜占用情况，通过搜索快速核对设备位置。'],['设备批量变更','导出设备信息，集中补充 IP、SN、功耗和责任人等字段，复核后批量回导更新。'],['机房扩容与资源规划','设置每排机柜数量，拖动调整大屏布局；通过多维利用率辅助判断扩容空间。'],['设备迁移与运维交接','在大屏中直接完成编辑、迁移和下架，结合位置履历、审计和审批记录降低交接遗漏。'],['多机房统一管理','按数据中心和机房分层查看资源，复制相似环境并在资源域之间移动机柜和机房。']],[1.5,5.05])
    d.add_heading('7. 产品特点与管理保障', 1)
    for s in ['编码灵活：数据中心、机房、机柜、设备和模板支持自动生成编码，也支持手工编码。','布局可靠：机柜使用固定格位，拖动只能交换或调整到有效位置，避免视觉和数据层面的重叠。','字段一致：设备台账、大屏编辑、详情、导入和导出以同一后台字段模型为准。','全系统统一视图：不区分设备正面和背面，减少维护歧义。','容量分析不虚报：缺少功耗、承重或 PDU 数据时展示“无法计算”及原因。','资源可持续扩展：不以固定业务数量上限限制数据中心、机房和机柜的持续建设。','安全可控：管理员权限、审批策略、LDAP 组映射和审计能力共同保障关键变更。','导入可复核：批量导入和机柜图回导先验证、后确认、再提交，支持识别新增、更新和潜在冲突。']: add_bullet(d,s)
    d.add_heading('8. 部署形态与技术基础', 1)
    d.add_paragraph('平台采用前后端分离架构，可通过 Docker Compose 快速部署；数据持久化使用 PostgreSQL。前端采用 Vue 3、TypeScript 和 Element Plus，后端采用 Go、Gin 和 GORM。LDAP 作为可选企业目录服务对接能力，可按组织现有身份体系接入。')
    d.add_heading('9. 对外产品信息', 1)
    table(d,['项目','内容'],[['产品名称','机柜管理工具'],['产品定位','数据中心基础资源可视化管理平台'],['正式产品版本',f'V{PRODUCT_VERSION}'],['适用对象','数据中心运维团队、机房管理员、基础设施管理人员、审计与管理人员'],['部署方式','Docker Compose / 私有化部署'],['数据存储','PostgreSQL'],['认证方式','本地账号、可选 LDAP 对接']],[1.5,5.05])
    d.add_heading('10. 结语', 1)
    d.add_paragraph('从资源台账到机房大屏，从设备上架到跨机房调整，机柜管理工具让数据中心资源管理更加清晰、准确、高效。平台将每一台设备、每一个 U 位和每一次变更连接起来，为数据中心的日常运维、容量规划和持续扩展提供统一基础。')
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('机柜管理工具\n让基础设施数据真正成为可视、可用、可管理的运营资产。'); r.bold=True; r.font.size=Pt(14); r.font.color.rgb=RGBColor(15,91,120)
    d.add_paragraph('截图说明：本文使用项目演示环境截图，截图中的名称、编号和时间为演示数据，仅用于功能说明，不代表任何客户生产环境。')
    d.save(DOCX_PATH)

# ---------- PDF ----------
def pdf_styles():
    font_path = Path(r'C:\Windows\Fonts\simhei.ttf')
    if font_path.exists(): pdfmetrics.registerFont(TTFont('SimHei', str(font_path)))
    else: pdfmetrics.registerFont(TTFont('SimHei', r'C:\Windows\Fonts\Deng.ttf'))
    ss=getSampleStyleSheet()
    return {
        'title': ParagraphStyle('p_title', parent=ss['Title'], fontName='SimHei', fontSize=28, leading=36, alignment=TA_CENTER, textColor=colors.HexColor('#0F5B78'), spaceAfter=12),
        'subtitle': ParagraphStyle('p_sub', parent=ss['Normal'], fontName='SimHei', fontSize=16, leading=23, alignment=TA_CENTER, textColor=colors.HexColor('#2E7D32'), spaceAfter=10),
        'tagline': ParagraphStyle('p_tag', parent=ss['Normal'], fontName='SimHei', fontSize=11, leading=18, alignment=TA_CENTER, textColor=colors.HexColor('#666666'), spaceAfter=15),
        'h1': ParagraphStyle('p_h1', parent=ss['Heading1'], fontName='SimHei', fontSize=17, leading=24, textColor=colors.HexColor('#0F5B78'), spaceBefore=10, spaceAfter=8),
        'h2': ParagraphStyle('p_h2', parent=ss['Heading2'], fontName='SimHei', fontSize=12.5, leading=18, textColor=colors.HexColor('#127C8C'), spaceBefore=8, spaceAfter=5),
        'body': ParagraphStyle('p_body', parent=ss['BodyText'], fontName='SimHei', fontSize=9.5, leading=16, spaceAfter=6),
        'small': ParagraphStyle('p_small', parent=ss['BodyText'], fontName='SimHei', fontSize=8, leading=12, textColor=colors.HexColor('#666666'), alignment=TA_CENTER),
        'bullet': ParagraphStyle('p_bullet', parent=ss['BodyText'], fontName='SimHei', fontSize=9.2, leading=15, leftIndent=14, firstLineIndent=-10, spaceAfter=3),
        'center': ParagraphStyle('p_center', parent=ss['BodyText'], fontName='SimHei', fontSize=12, leading=20, alignment=TA_CENTER, textColor=colors.HexColor('#0F5B78')),
    }

def ptxt(text, style):
    return Paragraph(text.replace('&','&amp;').replace('<','&lt;').replace('>','&gt;').replace('\n','<br/>'), style)

def rl_image(path, max_w=180*mm, max_h=105*mm):
    im=Image.open(path); w,h=im.size; scale=min(max_w/w,max_h/h); return RLImage(str(path), width=w*scale, height=h*scale, hAlign='CENTER')

def pdf_table(data, widths):
    t=Table(data, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([('FONTNAME',(0,0),(-1,-1),'SimHei'),('FONTSIZE',(0,0),(-1,-1),8),('LEADING',(0,0),(-1,-1),11),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#0F5B78')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),0.3,colors.HexColor('#B7D4DC')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#EDF6F8')]),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
    return t

def footer(canvas, doc):
    canvas.saveState(); canvas.setFont('SimHei',7.5); canvas.setFillColor(colors.HexColor('#888888')); canvas.drawCentredString(A4[0]/2, 10*mm, f'机柜管理工具 · 产品宣传文档 · V{PRODUCT_VERSION} · {TODAY}    |    {doc.page}'); canvas.restoreState()

def build_pdf():
    st=pdf_styles(); story=[]
    story += [Spacer(1,22*mm), ptxt('机柜管理工具',st['title']), ptxt('数据中心基础资源可视化管理平台',st['subtitle']), ptxt('让数据中心资源管理更清晰、更准确、更高效',st['tagline']), rl_image(SCREENSHOTS['screen_menu'],180*mm,83*mm), ptxt(f'正式产品版本：V{PRODUCT_VERSION}　|　文档版本：V{DOC_VERSION}　|　更新日期：{TODAY}',st['small']), PageBreak()]
    story += [ptxt('目录',st['h1'])]
    for x in ['1. 产品定位','2. 核心价值','3. 产品能力总览','4. 重点功能展示','5. 典型工作流程','6. 典型应用场景','7. 产品特点与管理保障','8. 部署形态与技术基础','9. 对外产品信息','10. 结语']: story.append(ptxt('• '+x,st['body']))
    story.append(PageBreak())
    story += [ptxt('1. 产品定位',st['h1']),ptxt('机柜管理工具面向数据中心、机房、机柜和设备的日常管理，将资源台账、U 位布局、设备信息、容量分析和运维操作统一在一个平台中。系统以机房大屏为核心入口，以结构化数据为基础，帮助运维团队从“靠记忆、靠表格、靠现场查找”转向“可视、可查、可操作、可追溯”。',st['body']),ptxt('平台适合新建机房初始化、存量资源盘点、设备批量上架、设备迁移、机柜扩容、运维交接和多机房统一管理等场景。',st['body']),ptxt('2. 核心价值',st['h1'])]
    for x in ['资源全景可视化：从数据中心到 U 位建立统一资源视图，机房状态一眼可见。','设备定位精确到 U 位：通过设备名称搜索，快速定位设备所在机柜和 U 位，减少现场查找时间。','变更操作形成闭环：设备上架、编辑、迁移、下架和归档均由业务流程管理，位置变化可追溯。','批量维护更高效：提供机柜、设备和机柜图的 Excel 模板、导入、导出及回导能力。','容量分析更可信：对 U 位、功耗、承重和 PDU 供电分别计算；条件不足时明确给出无法计算原因。','跨资源域调整更简单：支持数据中心、机房和机柜复制，以及跨机房、跨数据中心移动资源。','权限管理更稳妥：支持本地用户、角色权限、管理员审批和 LDAP 对接，满足团队化运维要求。']: story.append(ptxt('• '+x,st['bullet']))
    story += [ptxt('3. 产品能力总览',st['h1']),pdf_table([['能力域','主要能力'],['资源管理','数据中心、机房、机柜的增删改查、状态维护、复制与层级浏览'],['机房大屏','多机柜网格展示、固定格位、拖动交换位置、动态扩展和换行'],['U 位管理','空 U 位上架、有设备 U 位查看/编辑/迁移/下架，统一左键操作'],['设备台账','设备信息维护、自动编码、统一编辑、搜索定位、位置履历'],['模板与批量','机柜模板、机柜/设备批量导入、模板下载、设备信息导出后回导更新'],['机柜图','机柜图 Excel 导出与回导，支持导入验证和冲突检查'],['容量与供电','U 位、功耗、承重、PDU 插座与供电连接关系分析'],['权限治理','用户管理、角色权限、管理员审批、LDAP 用户和组映射、操作审计']],[34*mm,145*mm]),PageBreak()]
    story += [ptxt('4. 重点功能展示',st['h1']),ptxt('4.1 统一资源层级管理',st['h2']),ptxt('系统按照“数据中心 → 机房 → 机柜 → U 位 → 设备”组织资源。管理人员可以在同一页面浏览层级关系、查看资源数量和状态，并直接新增、编辑或进入更多操作。',st['body']),rl_image(SCREENSHOTS['resource'],180*mm,85*mm),ptxt('图 2 统一资源层级管理：以层级关系承载资源主数据',st['small']),ptxt('4.2 机房大屏与机柜可视化',st['h2']),ptxt('机房大屏以固定格位展示机柜，机柜不会相互重叠；支持通过拖动交换位置，并按照机房配置自动扩展和换行。管理人员可在大屏上直接新增、编辑、删除机柜和修改机柜状态。',st['body']),rl_image(SCREENSHOTS['room_status'],180*mm,84*mm),ptxt('图 3 机房大屏：多机柜状态、U 位占用和固定格位展示',st['small']),PageBreak()]
    story += [ptxt('4.3 U 位操作与设备生命周期',st['h2']),ptxt('平台统一使用左键点击 U 位：空 U 位可直接上架设备，有设备的 U 位可查看详情、编辑设备信息、调整/迁移位置或执行下架。设备位置由后端业务规则校验，避免重叠和重复占用。',st['body']),rl_image(SCREENSHOTS['screen_menu'],180*mm,84*mm),ptxt('图 4 U 位操作菜单：详情、调整/迁移、下架等操作集中呈现',st['small']),ptxt('4.4 设备搜索与快速定位',st['h2']),ptxt('输入设备名称并点击查询，即可在大屏中快速定位目标设备所在机柜及 U 位，适合设备巡检、故障处理和变更确认。',st['body']),rl_image(SCREENSHOTS['search'],180*mm,84*mm),ptxt('图 5 设备搜索：从名称快速定位机柜和 U 位',st['small']),PageBreak()]
    story += [ptxt('4.5 详情与容量分析',st['h2']),ptxt('机柜详情左侧展示机柜图，右侧集中展示机柜参数、U 位利用率、功耗利用率、承重利用率和 PDU 端口利用率。设备或机柜未补充必要数据时，系统不将缺失值默认为 0，而是输出无法计算的具体原因。',st['body']),rl_image(SCREENSHOTS['rack_detail'],180*mm,90*mm),ptxt('图 6 机柜详情与容量分析：参数、U 位、功耗和 PDU 信息一体化展示',st['small']),ptxt('4.6 统一设备编辑',st['h2']),ptxt('设备台账和机房大屏共用同一套设备编辑表单，字段以后台数据模型为准，确保新增、编辑、详情、批量导入和导出回导之间的数据一致。支持设备名称、设备编号、类型、厂商、型号、序列号、管理 IP、业务 IP、功耗、U 高度、负责人、用途和备注等信息。',st['body']),rl_image(SCREENSHOTS['device_editor'],165*mm,85*mm),ptxt('图 7 统一设备编辑表单：台账与大屏使用同一字段模型',st['small']),PageBreak()]
    story += [ptxt('4.7 模板、批量导入与设备信息回导',st['h2']),ptxt('平台提供清晰的 Excel 模板，区分必填字段和可选字段，并提供示例、字段说明和校验提示。普通机柜和设备批量导入不要求填写编码，编码为空时由系统自动生成。设备信息可先导出，再补充管理 IP、业务 IP、SN、功耗等字段后重新导入更新。机柜图也支持导出和回导，导入阶段先执行格式、设备识别、位置占用和重复数据验证，确认后再提交变更。',st['body']),rl_image(SCREENSHOTS['batch_editor'],180*mm,86*mm),ptxt('图 8 批量维护：通过模板和字段校验降低批量操作风险',st['small']),ptxt('4.8 复制与跨机房/跨数据中心移动',st['h2']),ptxt('支持复制数据中心、机房和机柜，用于快速构造相似资源环境；支持跨机房、跨数据中心移动机房或机柜，并在移动前校验目标层级、名称/编码冲突、U 位占用和设备位置关系。',st['body']),ptxt('4.9 PDU 与供电数据',st['h2']),ptxt('平台支持 PDU、插座和设备供电连接关系管理，插座类型可区分国标、欧标、10A 和 16A。PDU 信息属于可选配置，不会阻塞基础资源和设备管理；补充完整数据后，系统可进一步输出供电使用情况和利用率。',st['body'])]
    story += [ptxt('5. 典型工作流程',st['h1'])]
    for x in ['创建数据中心，编码可自动生成。','创建机房并设置机房基本信息和每排最多展示的机柜数量。','通过机柜模板创建机柜，按实际情况调整 U 位数、宽度、深度、高度等参数。','新增设备或通过 Excel 批量导入设备，编码可自动生成。','在大屏空 U 位上架设备，或通过迁移流程调整设备位置。','通过设备名称查询，快速定位设备所在机柜和 U 位。','查看机柜详情和利用率；补充功耗、承重或 PDU 数据后获得更完整分析。','进行设备编辑、迁移、下架、复制、跨机房移动并保留操作记录。']: story.append(ptxt('• '+x,st['bullet']))
    story += [ptxt('6. 典型应用场景',st['h1']),pdf_table([['场景','平台价值'],['新建数据中心初始化','统一创建数据中心、机房和机柜，使用模板和批量导入快速完成基础数据初始化。'],['存量机房资源盘点','导入现有设备信息，用大屏查看机柜占用情况，通过搜索快速核对设备位置。'],['设备批量变更','导出设备信息，集中补充 IP、SN、功耗和责任人等字段，复核后批量回导更新。'],['机房扩容与资源规划','设置每排机柜数量，拖动调整大屏布局；通过多维利用率辅助判断扩容空间。'],['设备迁移与运维交接','在大屏中直接完成编辑、迁移和下架，结合位置履历、审计和审批记录降低交接遗漏。'],['多机房统一管理','按数据中心和机房分层查看资源，复制相似环境并在资源域之间移动机柜和机房。']],[38*mm,141*mm]),ptxt('7. 产品特点与管理保障',st['h1'])]
    for x in ['编码灵活：数据中心、机房、机柜、设备和模板支持自动生成编码，也支持手工编码。','布局可靠：机柜使用固定格位，拖动只能交换或调整到有效位置，避免视觉和数据层面的重叠。','字段一致：设备台账、大屏编辑、详情、导入和导出以同一后台字段模型为准。','全系统统一视图：不区分设备正面和背面，减少维护歧义。','容量分析不虚报：缺少功耗、承重或 PDU 数据时展示“无法计算”及原因。','资源可持续扩展：不以固定业务数量上限限制数据中心、机房和机柜的持续建设。','安全可控：管理员权限、审批策略、LDAP 组映射和审计能力共同保障关键变更。','导入可复核：批量导入和机柜图回导先验证、后确认、再提交，支持识别新增、更新和潜在冲突。']: story.append(ptxt('• '+x,st['bullet']))
    story += [ptxt('8. 部署形态与技术基础',st['h1']),ptxt('平台采用前后端分离架构，可通过 Docker Compose 快速部署；数据持久化使用 PostgreSQL。前端采用 Vue 3、TypeScript 和 Element Plus，后端采用 Go、Gin 和 GORM。LDAP 作为可选企业目录服务对接能力，可按组织现有身份体系接入。',st['body']),ptxt('9. 对外产品信息',st['h1']),pdf_table([['项目','内容'],['产品名称','机柜管理工具'],['产品定位','数据中心基础资源可视化管理平台'],['正式产品版本',f'V{PRODUCT_VERSION}'],['适用对象','数据中心运维团队、机房管理员、基础设施管理人员、审计与管理人员'],['部署方式','Docker Compose / 私有化部署'],['数据存储','PostgreSQL'],['认证方式','本地账号、可选 LDAP 对接']],[38*mm,141*mm]),ptxt('10. 结语',st['h1']),ptxt('从资源台账到机房大屏，从设备上架到跨机房调整，机柜管理工具让数据中心资源管理更加清晰、准确、高效。平台将每一台设备、每一个 U 位和每一次变更连接起来，为数据中心的日常运维、容量规划和持续扩展提供统一基础。',st['body']),Spacer(1,8*mm),ptxt('机柜管理工具<br/>让基础设施数据真正成为可视、可用、可管理的运营资产。',st['center']),Spacer(1,8*mm),ptxt('截图说明：本文使用项目演示环境截图，截图中的名称、编号和时间为演示数据，仅用于功能说明，不代表任何客户生产环境。',st['small'])]
    pdf=SimpleDocTemplate(str(PDF_PATH),pagesize=A4,rightMargin=15*mm,leftMargin=15*mm,topMargin=13*mm,bottomMargin=17*mm,title='机柜管理工具产品宣传文档',author='Cabinet Management Tool')
    pdf.build(story,onFirstPage=footer,onLaterPages=footer)

# ---------- run ----------
def main():
    build_docx(); build_pdf()
    md=md_content(); MD_PATH.write_text(md,encoding='utf-8'); PROJECT_MD_PATH.write_text(md,encoding='utf-8')
    print(DOCX_PATH); print(PDF_PATH); print(MD_PATH); print(PROJECT_MD_PATH)

if __name__=='__main__': main()
